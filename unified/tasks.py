"""Server-authoritative task lifecycle with ownership and Mongo compare-and-swap."""
from __future__ import annotations
import asyncio,copy,csv,hashlib,io,json,re,uuid
from datetime import datetime,timezone
from .tools import execute
class Conflict(ValueError):pass
class Missing(PermissionError):pass
def digest(x):return hashlib.sha256(json.dumps(x,sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
def now():return datetime.now(timezone.utc).isoformat()
class MongoTasks:
 def __init__(self,db):self.collection=db['workbench_tasks']
 async def insert(self,row):await self.collection.insert_one(row)
 async def get(self,id,owner):return await self.collection.find_one({'_id':id,'owner':owner})
 async def change(self,id,owner,revision,row):
  res=await self.collection.replace_one({'_id':id,'owner':owner,'revision':revision},row)
  if res.modified_count!=1:raise Conflict('任务版本已变化，请刷新后重试')
 async def list(self,owner):return await self.collection.find({'owner':owner}).sort('updated_at',-1).limit(30).to_list(length=30)
class TaskService:
 def __init__(self,repo,runner=execute):self.repo=repo;self.runner=runner
 async def read(self,id,owner):
  row=await self.repo.get(id,owner)
  if not row:raise Missing('任务不存在或不属于当前用户')
  return copy.deepcopy(row)
 async def create(self,owner,tool,values,sources=None,use_examples=False):
  if len(json.dumps({'values':values,'sources':sources},ensure_ascii=False).encode())>1_100_000:raise ValueError('单任务最多1MB输入数据')
  row={'_id':uuid.uuid4().hex,'owner':owner,'tool':tool,'values':values,'sources':sources,'use_examples':use_examples,'revision':1,'status':'draft','result':None,'approval':None,'created_at':now(),'updated_at':now()}
  await self.repo.insert(row);return row
 async def edit(self,id,owner,revision,values,sources,use_examples):
  if len(json.dumps({'values':values,'sources':sources},ensure_ascii=False).encode())>1_100_000:raise ValueError('单任务最多1MB输入数据')
  row=await self.read(id,owner)
  if row['revision']!=revision:raise Conflict('旧输入版本')
  row.update(values=values,sources=sources,use_examples=use_examples,revision=revision+1,status='draft',result=None,approval=None,updated_at=now())
  await self.repo.change(id,owner,revision,row);return row
 async def run(self,id,owner,revision):
  row=await self.read(id,owner)
  if row['revision']!=revision or row['status'] not in {'draft','failed'}:raise Conflict('任务已运行或版本变化')
  row.update(status='running',revision=revision+1,updated_at=now());await self.repo.change(id,owner,revision,row);running=row['revision']
  try:result=await asyncio.to_thread(self.runner,row['tool'],row['values'],examples=row['use_examples'],sources=row['sources'])
  except Exception as e:row.update(status='failed',error=(str(e)[:500] if isinstance(e,(ValueError,KeyError,TypeError)) else type(e).__name__),revision=running+1,updated_at=now())
  else:row.update(status='needs_review',result=result,result_hash=digest(result),revision=running+1,updated_at=now())
  await self.repo.change(id,owner,running,row);return row
 async def approve(self,id,owner,revision):
  row=await self.read(id,owner)
  if row['revision']!=revision or row['status']!='needs_review':raise Conflict('只能确认当前待复核结果')
  row.update(status='approved',approval={'actor':owner,'result_hash':row['result_hash'],'at':now()},revision=revision+1,updated_at=now())
  await self.repo.change(id,owner,revision,row);return row
 async def export(self,id,owner,revision,fmt):
  row=await self.read(id,owner)
  if row['revision']!=revision or row['status']!='approved' or row['approval']['result_hash']!=digest(row['result']):raise Conflict('请先确认当前结果，修改后须重新复核')
  content,kind=await asyncio.to_thread(render_export,row,fmt);latest=await self.read(id,owner)
  if latest['revision']!=revision:raise Conflict('文件生成期间输入改变，停止旧导出')
  return content,kind

def render_export(row,fmt):
 result=row['result']['result'];rows=result.get('rows',[])
 if fmt=='json':return json.dumps(row,ensure_ascii=False,indent=2).encode(),'application/json'
 if fmt=='csv':
  buf=io.StringIO();heads=list(dict.fromkeys(k for x in rows for k in x))
  if not heads:raise ValueError('当前工具没有表格可导出，请选JSON')
  writer=csv.writer(buf)
  def cell(v):
   text=json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else '' if v is None else str(v)
   if text.lstrip()[:1] in '=+@-' and not re.fullmatch(r'[+-]?\d+(\.\d+)?',text):text="'"+text
   return text
  writer.writerow([cell(k) for k in heads])
  for x in rows:writer.writerow([cell(x.get(k)) for k in heads])
  return ('\ufeff'+buf.getvalue()).encode(),'text/csv; charset=utf-8'
 if fmt=='docx':
  from docx import Document
  d=Document();d.add_heading('金枢｜待业务复核的任务材料',0)
  d.add_paragraph('工具：'+row['tool']);d.add_paragraph('任务版本：'+str(row['revision']))
  d.add_paragraph('数据模式：'+row['result']['input_mode']);d.add_paragraph('确认人：'+row['approval']['actor'])
  d.add_paragraph('产品名称：'+str(result.get('product_name','未提供')))
  heads=list(dict.fromkeys(k for x in rows for k in x))
  if len(heads)>12 or len(rows)>500:raise ValueError('Word最多500行、12列，请用CSV/JSON导出完整数据')
  if heads:
   table=d.add_table(rows=1,cols=len(heads));table.style='Table Grid'
   for c,h in zip(table.rows[0].cells,heads):c.text=h
   for x in rows:
    for c,h in zip(table.add_row().cells,heads):c.text=str(x.get(h,''))
  d.add_paragraph(result.get('note',''));d.add_paragraph('本文件不代表发行、付款、账户处置或监管结论。')
  b=io.BytesIO();d.save(b);return b.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
 raise ValueError('只支持json、csv或docx')
