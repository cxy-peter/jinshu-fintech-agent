"""Build a private, portable page-text pack. Originals never leave this computer.

Run locally with PyMuPDF and python-docx. Not a cloud build dependency. Output
belongs in private_corpus (git/deploy ignored), never in the public source tree.
Native text extraction does not assert that financial table layout is correct.
"""
import argparse
import hashlib
import json
from pathlib import Path


def build(inputs, destination):
    destination = Path(destination).resolve()
    if 'private_corpus' not in destination.parts:
        raise ValueError('Output must be inside a private_corpus directory')
    destination.parent.mkdir(parents=True, exist_ok=True)
    docs, failures, seen, duplicates = [], [], set(), 0
    files = sorted({f.resolve() for root in inputs for f in
                    (Path(root).rglob('*') if Path(root).is_dir() else [Path(root)])
                    if f.suffix.lower() in {'.pdf', '.docx', '.txt', '.md'}})
    for f in files:
        sha = hashlib.sha256(f.read_bytes()).hexdigest()
        if sha in seen:
            duplicates += 1
            continue
        seen.add(sha)
        try:
            if f.suffix.lower() == '.pdf':
                import fitz
                with fitz.open(f) as pdf:
                    if pdf.is_encrypted:
                        raise ValueError('encrypted')
                    pages = [{'page': i + 1, 'text': p.get_text(sort=True)}
                             for i, p in enumerate(pdf)]
            elif f.suffix.lower() == '.docx':
                from docx import Document
                d = Document(f)
                body = '\n'.join(p.text for p in d.paragraphs)
                body += '\n' + '\n'.join('| ' + ' | '.join(c.text for c in r.cells) + ' |'
                                           for t in d.tables for r in t.rows)
                pages = [{'page': 1, 'text': body}]
            else:
                pages = [{'page': 1, 'text': f.read_text(encoding='utf-8-sig')}]
            if not any(p['text'].strip() for p in pages):
                raise ValueError('no_native_text')
            docs.append(dict(id='personal_' + sha[:24], title=f.name, sha256=sha,
                             status='active', flow='all', version=1, pages=pages,
                             source_kind='personal_reference', synthetic=False,
                             authority='个人研究与学习资料；不自动视为现行制度',
                             parser='native_text_not_table_layout_verified'))
        except Exception as exc:
            failures.append({'file': f.name, 'error': type(exc).__name__})
    result = dict(format='jinshu-personal-v1', documents=docs,
                  scope='private_local_reference_only', shared_review=False)
    data = json.dumps(result, ensure_ascii=False, separators=(',', ':'))
    if len(data.encode()) > 100_000_000:
        raise ValueError('Pack exceeds 100 MB; split the input folders')
    temporary = destination.with_suffix('.tmp')
    temporary.write_text(data, encoding='utf-8')
    temporary.replace(destination)
    report = dict(files=len(files), documents=len(docs), duplicates=duplicates,
                  pages=sum(len(d['pages']) for d in docs), bytes=destination.stat().st_size,
                  sha256=hashlib.sha256(destination.read_bytes()).hexdigest(),
                  failures=failures, originals_changed=False, uploaded=False)
    destination.with_suffix('.manifest.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('inputs', nargs='+')
    p.add_argument('--out', required=True)
    args = p.parse_args()
    print(json.dumps(build(args.inputs, args.out), ensure_ascii=False, indent=2))
