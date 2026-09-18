"""Registered Python skills based on prior reporting/issuance methods.
No user-provided Python is executed. New functions are registered and reviewed in code.
"""
from __future__ import annotations
import hashlib,json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from . import ROOT,tools
from pydantic import BaseModel,Field,ConfigDict
from typing import Literal
from datetime import date

class Params(BaseModel):
    model_config=ConfigDict(extra="forbid")
class Benchmark(Params):
    product:str="SIM0001"
    start:date=date(2026,6,1)
    end:date=date(2026,8,30)
class Schedule(Params):
    start:date=date(2026,9,25)
    term_days:int=Field(default=90,ge=1,le=3650)
    frequency_days:Literal[7,14]=7
    count:int=Field(default=3,ge=1,le=20)
class Material(Schedule):
    product_name:str=Field(default="模拟稳健产品A",min_length=1,max_length=80)
SCHEMAS={"wealth_benchmark":Benchmark,"issuance":Schedule,"weekly_report":Params,"material_fill":Material}

REGISTRY={
 'wealth_benchmark':{'entrypoint':'jinshu.tools:wealth_benchmark','input':'product,start,end','output':'peer rows, common period, median, drawdown'},
 'issuance':{'entrypoint':'jinshu.tools:issuance','input':'start,term_days,frequency_days,count','output':'schedule rows, calendar version'},
 'weekly_report':{'entrypoint':'jinshu.tools:weekly_report','input':'registered mock CSV','output':'normalized rows and missing-data issues'},
 'material_fill':{'entrypoint':'jinshu.python_skills:material_fill','input':'product_name,start,term_days,frequency_days,count','output':'Word draft and schedule'},
}

def catalog():return [{'skill':key,**value,'schema':SCHEMAS[key].model_json_schema(),'mode':'registered_python'} for key,value in REGISTRY.items()]

def material_fill(params):
    schedule=tools.issuance(params)
    product_name=str(params.get('product_name','模拟稳健产品A'))[:80]
    payload={'product_name':product_name,'schedule':schedule}
    digest=hashlib.sha256(json.dumps(payload,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:20]
    directory=ROOT/'workspace'/'outputs';directory.mkdir(parents=True,exist_ok=True)
    path=directory/f'issuance-{digest}.docx'
    if not path.exists():
        doc=Document();doc.add_heading('模拟理财产品发行材料草稿',0)
        doc.add_paragraph('个人项目合成样例；不是公司正式发行文件。')
        doc.add_paragraph('产品名称：'+product_name)
        doc.add_paragraph('日历版本：'+schedule['material_preview']['日历版本'])
        headers=['批次','募集结束日','成立日','到期日','实际自然日']
        t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
        for c,h in zip(t.rows[0].cells,headers):c.text=h
        for r in schedule['rows']:
            for c,h in zip(t.add_row().cells,headers):c.text=str(r[h])
        doc.add_heading('待复核事项',1)
        for item in schedule['material_preview']['需人工确认']:doc.add_paragraph(item)
        doc.add_paragraph('处理路径：参数校验 → 调用现有排期函数 → 映射模板字段 → 输出草稿。')
        doc.save(path)
    return {'rows':schedule['rows'],'output_file':path.name,'download_path':'/api/outputs/'+path.name,
            'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'note':'模板草稿；重复相同输入复用同一文件，不自动发送或批准发行。'}

def execute(name,params):
    if name not in REGISTRY:raise ValueError('Python skill未登记')
    params=SCHEMAS[name].model_validate(params).model_dump(mode='json')
    if name=='material_fill':
        result={'tool':name,'synthetic':True,'result':material_fill(params),'source_files':[{'file':'calendar.json','sha256':hashlib.sha256((tools.DATA/'calendar.json').read_bytes()).hexdigest()}]}
    else:result=tools.run(name,params)
    return result|{'python_skill':REGISTRY[name]}
